"""
Generates the PowerPoint Presentation (.pptx) strictly following the hackathon template.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(5.625)

PRIMARY_COLOR = RGBColor(27, 94, 32)
TEXT_COLOR = RGBColor(33, 33, 33)

def add_header(slide, title_text):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_COLOR
    p.alignment = PP_ALIGN.CENTER

slide_layout = prs.slide_layouts[6] # blank

# -------------------------------------------------------------
# SLIDE 1: Title Page
# -------------------------------------------------------------
slide1 = prs.slides.add_slide(slide_layout)
box1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(8.0), Inches(3.8))
tf1 = box1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "SUPERNOVA 2.0 - STELLAR HACKATHON 2025"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = PRIMARY_COLOR
p.alignment = PP_ALIGN.CENTER

p_sub = tf1.add_paragraph()
p_sub.text = "TITLE PAGE\n"
p_sub.font.size = Pt(14)
p_sub.font.bold = True
p_sub.font.color.rgb = TEXT_COLOR
p_sub.alignment = PP_ALIGN.CENTER

bullets1 = [
    "• Problem Statement ID – [Enter PS ID]",
    "• Problem Statement Title – Market-Linked Crop Planning, Rural Advisory and Agricultural Operations Platform",
    "• Theme – Smart Agriculture / Rural Development",
    "• PS Category – Software",
    "• Team ID – [Enter Team ID]",
    "• Team Name (Registered) – [Enter Team Name]"
]
for b in bullets1:
    p_b = tf1.add_paragraph()
    p_b.text = b
    p_b.font.size = Pt(12)
    p_b.font.color.rgb = TEXT_COLOR

# -------------------------------------------------------------
# SLIDE 2: Idea Title & Proposed Solution
# -------------------------------------------------------------
slide2 = prs.slides.add_slide(slide_layout)
add_header(slide2, "KrishiSetu: IDEA TITLE")

box2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(3.8))
tf2 = box2.text_frame
tf2.word_wrap = True

p_heading = tf2.paragraphs[0]
p_heading.text = "❖ Proposed Solution (Describe your Idea/Solution/Prototype)"
p_heading.font.size = Pt(15)
p_heading.font.bold = True
p_heading.font.color.rgb = PRIMARY_COLOR

s2_bullets = [
    "• Detailed explanation of the proposed solution:\n  - Full-stack Python/Streamlit platform connecting agricultural supply to mandi demand.\n  - Secure SQLite authentication with multi-persona roles (Farmers, FPOs, Buyers, Agronomists).",
    "• How it addresses the problem:\n  - Solves price volatility and crop mismatch via ROI crop planning, real-time weather risk advisories, and direct pre-harvest buyer contracts.",
    "• Innovation and uniqueness of the solution:\n  - Multilingual support (English & Marathi) with 'Krishi-Sarthi' WhatsApp/SMS voice bot simulator for low-bandwidth rural accessibility."
]
for b in s2_bullets:
    p = tf2.add_paragraph()
    p.text = b
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_COLOR

# -------------------------------------------------------------
# SLIDE 3: Technical Approach
# -------------------------------------------------------------
slide3 = prs.slides.add_slide(slide_layout)
add_header(slide3, "TECHNICAL APPROACH")

box3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.4), Inches(3.6))
tf3 = box3.text_frame
tf3.word_wrap = True

s3_bullets = [
    "• Technologies to be used (programming languages, frameworks, hardware):\n  - Frontend & Dashboard: Python Streamlit\n  - Backend & Database: Python, SQLite3 (Auth, Chat History, Escrow Transactions)\n  - AI & Machine Learning: Google Gemini 1.5 Flash Vision API (Leaf Pathology) & Scikit-Learn\n  - External APIs: Open-Meteo REST API (Hyper-local weather forecasting)",
    "• Methodology and process for implementation (Flow Charts / Working Prototype):\n  - Modular service architecture (crop_planner.py, weather_service.py, crop_doctor.py, marketplace.py)\n  - Offline-first PWA caching layer for continuous rural uptime during connectivity drops."
]
for b in s3_bullets:
    p = tf3.add_paragraph()
    p.text = b
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_COLOR

# -------------------------------------------------------------
# SLIDE 4: Feasibility and Viability
# -------------------------------------------------------------
slide4 = prs.slides.add_slide(slide_layout)
add_header(slide4, "FEASIBILITY AND VIABILITY")

box4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.4), Inches(3.6))
tf4 = box4.text_frame
tf4.word_wrap = True

s4_bullets = [
    "• Analysis of the feasibility of the idea:\n  - Built on lightweight, open-source Python stacks operable on standard cloud servers or local devices with minimal compute overhead.",
    "• Potential challenges and risks:\n  - Unstable rural internet connectivity and digital literacy barriers for smallholder farmers.",
    "• Strategies for overcoming these challenges:\n  - Localized Marathi vernacular support, zero-learning-curve WhatsApp/SMS bot simulation ('Krishi-Sarthi'), and automated offline caching."
]
for b in s4_bullets:
    p = tf4.add_paragraph()
    p.text = b
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_COLOR

# -------------------------------------------------------------
# SLIDE 5: Impact and Benefits
# -------------------------------------------------------------
slide5 = prs.slides.add_slide(slide_layout)
add_header(slide5, "IMPACT AND BENEFITS")

box5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.4), Inches(3.6))
tf5 = box5.text_frame
tf5.word_wrap = True

s5_bullets = [
    "• Potential impact on the target audience:\n  - Empowers over 15+ million smallholder farmers across Maharashtra's 36 districts with real-time market linkages and agronomic guidance.",
    "• Benefits of the solution (social, economic, environmental, etc.):\n  - Social: Bridges the digital divide via vernacular voice/text assistance.\n  - Economic: Eliminates middlemen through direct pre-harvest contracts, increasing profit margins by 20-30%.\n  - Environmental: Promotes precise fertilizer usage and prevents chemical runoff via weather risk alerts."
]
for b in s5_bullets:
    p = tf5.add_paragraph()
    p.text = b
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_COLOR

# -------------------------------------------------------------
# SLIDE 6: Research and References
# -------------------------------------------------------------
slide6 = prs.slides.add_slide(slide_layout)
add_header(slide6, "RESEARCH AND REFERENCES")

box6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.4), Inches(3.6))
tf6 = box6.text_frame
tf6.word_wrap = True

s6_bullets = [
    "• Details / Links of the reference and research work:\n  - Open-Meteo Weather API Documentation: https://open-meteo.com/\n  - Google Generative AI Python SDK: https://ai.google.dev/\n  - Indian Council of Agricultural Research (ICAR) Baseline Agronomic Norms\n  - Maharashtra Mandi Price Trends & Agmarknet Datasets\n  - Local Workspace Repository: KrishiSetu Platform Architecture"
]
for b in s6_bullets:
    p = tf6.add_paragraph()
    p.text = b
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_COLOR

prs.save("KrishiSetu_Hackathon_Presentation.pptx")
print("Updated presentation saved successfully as KrishiSetu_Hackathon_Presentation.pptx")
